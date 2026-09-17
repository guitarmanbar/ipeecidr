#!/usr/bin/env python3
"""
IPv4 Network Framework
- Загружает модель сети из Excel (CIDR, first_ip, last_ip, short_name, description).
- Функция 1: для списка IP находит все подсети, в которые они входят.
- Функция 2: агрегирует IP в минимальный набор CIDR, точно покрывающий только эти IP.
- Интерактивный REPL.
"""

import os
import sys
import ipaddress
import traceback
from openpyxl import Workbook, load_workbook

# ----------------------------------------------------------------------
# Модель сети
# ----------------------------------------------------------------------
class NetworkModel:
    def __init__(self, filepath):
        self.subnets = []
        self.load(filepath)

    def load(self, filepath):
        """Загружает модель сети из Excel. Гибко определяет колонки по заголовкам."""
        wb = load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Excel-файл пуст")

        headers = [str(h).strip().lower() if h else '' for h in rows[0]]
        cidr_col = name_col = desc_col = None

        for i, h in enumerate(headers):
            if 'cidr' in h:
                cidr_col = i
            elif 'name' in h or 'название' in h:
                name_col = i
            elif 'desc' in h or 'описание' in h:
                desc_col = i

        if cidr_col is None:
            raise ValueError("Не найдена колонка с CIDR")

        for row in rows[1:]:
            if not row or row[cidr_col] is None:
                continue
            cidr_str = str(row[cidr_col]).strip()
            try:
                network = ipaddress.IPv4Network(cidr_str, strict=False)
            except Exception as e:
                raise ValueError(f"Некорректный CIDR '{cidr_str}': {e}")

            name = str(row[name_col]).strip() if name_col is not None and row[name_col] is not None else ''
            desc = str(row[desc_col]).strip() if desc_col is not None and row[desc_col] is not None else ''

            self.subnets.append({
                'cidr': str(network),
                'network': network,
                'name': name,
                'description': desc
            })

    def find_subnets(self, ip_str):
        """Возвращает список подсетей, в которые входит IP."""
        ip = ipaddress.IPv4Address(ip_str)
        return [s for s in self.subnets if ip in s['network']]


# ----------------------------------------------------------------------
# Чтение IP из файла (Excel или текстовый)
# ----------------------------------------------------------------------
def read_ips(filepath):
    """Читает IP из Excel (первая колонка) или текстового файла (по одному в строке).
       Убирает пробелы, удаляет дубликаты, валидирует."""
    ext = os.path.splitext(filepath)[1].lower()
    raw_ips = []

    if ext == '.xlsx':
        wb = load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            if row and row[0] is not None:
                val = str(row[0]).strip()
                if val:
                    raw_ips.append(val)
    else:
        with open(filepath, 'r', encoding='utf-8') as f:
            for line in f:
                val = line.strip()
                if val:
                    raw_ips.append(val)

    unique_ips = []
    seen = set()
    for ip_str in raw_ips:
        try:
            ip = ipaddress.IPv4Address(ip_str)
            if ip not in seen:
                seen.add(ip)
                unique_ips.append(ip)
        except Exception as e:
            raise ValueError(f"Некорректный IP-адрес '{ip_str}': {e}")

    return unique_ips


# ----------------------------------------------------------------------
# Функция 1: поиск подсетей для IP
# ----------------------------------------------------------------------
def lookup_ips(network_model, ips, output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Lookup Results"
    ws.append(["IP-address", "CIDR", "Subnet Name", "Description", "Duplicate"])

    for ip in ips:
        matches = network_model.find_subnets(str(ip))
        if not matches:
            ws.append([str(ip), "NOT_FOUND", "", "", False])
        else:
            duplicate = len(matches) > 1
            for match in matches:
                ws.append([str(ip), match['cidr'], match['name'], match['description'], duplicate])

    wb.save(output_file)
    print(f"Результаты поиска сохранены в {output_file}")


# ----------------------------------------------------------------------
# Функция 2: агрегация IP в минимальный набор CIDR
# ----------------------------------------------------------------------
def aggregate_ips(ips, output_file):
    """Строит минимальный набор CIDR, точно покрывающий только переданные IP."""
    if not ips:
        raise ValueError("Нет IP-адресов для агрегации")

    sorted_ips = sorted(ips, key=lambda x: int(x))
    ranges = []
    start = end = sorted_ips[0]

    for ip in sorted_ips[1:]:
        if int(ip) == int(end) + 1:
            end = ip
        else:
            ranges.append((start, end))
            start = end = ip
    ranges.append((start, end))

    cidrs = []
    for start, end in ranges:
        for cidr in ipaddress.summarize_address_range(start, end):
            cidrs.append(cidr)

    wb = Workbook()
    ws = wb.active
    ws.title = "Aggregated CIDRs"
    ws.append(["CIDR", "Netmask", "Number of IPs"])
    for cidr in cidrs:
        ws.append([str(cidr), str(cidr.netmask), cidr.num_addresses])

    wb.save(output_file)
    print(f"Агрегированные CIDR сохранены в {output_file}")


# ----------------------------------------------------------------------
# Интерактивный REPL
# ----------------------------------------------------------------------
def repl():
    network_model = None
    print("IPv4 Network Framework REPL")
    print("Доступные команды:")
    print("  load <network.xlsx>                     - загрузить модель сети")
    print("  lookup <ips_file> [output.xlsx]         - найти подсети для IP")
    print("  aggregate <ips_file> [output.xlsx]      - агрегировать IP в CIDR")
    print("  help                                    - показать справку")
    print("  exit                                    - выход")

    while True:
        try:
            cmd = input("> ").strip()
            if not cmd:
                continue
            parts = cmd.split()
            action = parts[0].lower()

            if action == 'exit':
                break
            elif action == 'help':
                print("Команды: load, lookup, aggregate, help, exit")
            elif action == 'load':
                if len(parts) < 2:
                    print("Использование: load <network.xlsx>")
                    continue
                network_model = NetworkModel(parts[1])
                print(f"Загружено подсетей: {len(network_model.subnets)}")
            elif action == 'lookup':
                if network_model is None:
                    print("Сначала загрузите модель сети командой 'load'")
                    continue
                if len(parts) < 2:
                    print("Использование: lookup <ips_file> [output.xlsx]")
                    continue
                ips_file = parts[1]
                output_file = parts[2] if len(parts) > 2 else 'lookup_output.xlsx'
                ips = read_ips(ips_file)
                print(f"Прочитано уникальных IP: {len(ips)}")
                lookup_ips(network_model, ips, output_file)
            elif action == 'aggregate':
                if len(parts) < 2:
                    print("Использование: aggregate <ips_file> [output.xlsx]")
                    continue
                ips_file = parts[1]
                output_file = parts[2] if len(parts) > 2 else 'aggregate_output.xlsx'
                ips = read_ips(ips_file)
                print(f"Прочитано уникальных IP: {len(ips)}")
                aggregate_ips(ips, output_file)
            else:
                print(f"Неизвестная команда: {action}")
        except Exception as e:
            traceback.print_exc()
            print(f"Ошибка: {e}")


if __name__ == '__main__':
    repl()
